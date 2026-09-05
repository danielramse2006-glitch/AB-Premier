import base64
import os
import secrets
import sqlite3
import sys
import threading
import time
import csv
import socket
from io import StringIO, BytesIO
from datetime import datetime
from pathlib import Path

import webview
from flask import Flask, jsonify, render_template, request, session
from waitress import serve
from werkzeug.security import check_password_hash, generate_password_hash


APP_NAME = "AB Premiere"
HOST = "127.0.0.1"
PORT = 5001


def resource_path(*parts):
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base.joinpath(*parts)


def data_dir():
    override = os.getenv("AB_PREMIERE_DATA_DIR")
    if override:
        root = Path(override)
    elif getattr(sys, "frozen", False):
        root = Path(os.getenv("LOCALAPPDATA", Path.home())) / "AB Premiere"
    else:
        root = Path(__file__).resolve().parent / "data"
    root.mkdir(parents=True, exist_ok=True)
    (root / "photos").mkdir(exist_ok=True)
    return root


DB_PATH = data_dir() / "ab_premiere.db"
app = Flask(__name__, template_folder=str(resource_path("templates")), static_folder=str(resource_path("static")))
def persistent_secret():
    secret_file = data_dir() / ".secret"
    if not secret_file.exists():
        secret_file.write_text(secrets.token_hex(32), encoding="utf-8")
    return secret_file.read_text(encoding="utf-8").strip()


app.secret_key = os.getenv("AB_PREMIERE_SECRET", persistent_secret())
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Strict")


def db():
    conn = sqlite3.connect(DB_PATH, timeout=15)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA busy_timeout=15000")
    return conn


def init_db():
    with db() as conn:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS admins (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE COLLATE NOCASE,
            password_hash TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS clients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT NOT NULL UNIQUE COLLATE NOCASE,
            full_name TEXT NOT NULL,
            phone TEXT DEFAULT '',
            photo TEXT DEFAULT '',
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS visits (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client_id INTEGER NOT NULL,
            visited_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(client_id) REFERENCES clients(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS idx_visits_client ON visits(client_id);
        CREATE INDEX IF NOT EXISTS idx_visits_date ON visits(visited_at);
        CREATE TABLE IF NOT EXISTS audit_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT, admin_id INTEGER,
            action TEXT NOT NULL, entity TEXT NOT NULL, entity_id INTEGER,
            details TEXT DEFAULT '', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        """)
        admin_columns = {r[1] for r in conn.execute("PRAGMA table_info(admins)")}
        for name, definition in {"full_name":"TEXT DEFAULT ''","role":"TEXT NOT NULL DEFAULT 'admin'","last_login":"TEXT","failed_attempts":"INTEGER NOT NULL DEFAULT 0","locked_until":"TEXT"}.items():
            if name not in admin_columns: conn.execute(f"ALTER TABLE admins ADD COLUMN {name} {definition}")
        visit_columns = {r[1] for r in conn.execute("PRAGMA table_info(visits)")}
        for name, definition in {"notes":"TEXT DEFAULT ''","created_by":"INTEGER","source":"TEXT DEFAULT 'portal'"}.items():
            if name not in visit_columns: conn.execute(f"ALTER TABLE visits ADD COLUMN {name} {definition}")
        if not conn.execute("SELECT 1 FROM admins LIMIT 1").fetchone():
            conn.execute("INSERT INTO admins(username,password_hash) VALUES(?,?)",
                         ("admin", generate_password_hash("ABPremier2026")))
        conn.execute("INSERT OR IGNORE INTO settings(key,value) VALUES('release_code',?)",
                     (generate_password_hash("ABREYNOSA"),))
        for key, value in (("tier_frequent","5"),("tier_premiere","10"),("business_name","AB PREMIERE")):
            conn.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)",(key,value))
        if not conn.execute("SELECT 1 FROM clients LIMIT 1").fetchone():
            conn.execute("INSERT INTO clients(code,full_name,phone) VALUES(?,?,?)",("2026","Cliente Demostración","899-000-0000"))
        # Migrar códigos anteriores AB-0000-VIP al formato numérico de 4 dígitos.
        for old in conn.execute("SELECT id,code FROM clients WHERE code LIKE 'AB-____-VIP'").fetchall():
            digits = old["code"][3:7]
            if digits.isdigit() and not conn.execute("SELECT 1 FROM clients WHERE code=? AND id<>?",(digits,old["id"])).fetchone():
                conn.execute("UPDATE clients SET code=? WHERE id=?",(digits,old["id"]))


def json_error(message, status=400):
    return jsonify(ok=False, message=message), status


def require_admin():
    return bool(session.get("admin_id"))


def audit(action, entity, entity_id=None, details=""):
    with db() as conn:
        conn.execute("INSERT INTO audit_log(admin_id,action,entity,entity_id,details) VALUES(?,?,?,?,?)",
                     (session.get("admin_id"),action,entity,entity_id,str(details)[:1000]))


def save_photo(data_url, code):
    if not data_url or "," not in data_url:
        return ""
    header, encoded = data_url.split(",", 1)
    if not header.startswith("data:image/"):
        raise ValueError("Formato de foto inválido")
    raw = base64.b64decode(encoded, validate=True)
    if len(raw) > 3 * 1024 * 1024:
        raise ValueError("La foto supera 3 MB")
    filename = f"{code}-{secrets.token_hex(4)}.jpg"
    (data_dir() / "photos" / filename).write_bytes(raw)
    return filename


def client_payload(row, visits):
    count = int(visits)
    tier = "Premiere ⭐" if count >= 10 else "Frecuente" if count >= 5 else "Socio AB"
    photo = f"/api/photos/{row['photo']}" if row["photo"] else ""
    return {"id": row["id"], "code": row["code"], "name": row["full_name"],
            "phone": row["phone"], "photo": photo, "visits": count, "tier": tier,
            "active": bool(row["active"]), "created_at": row["created_at"]}


@app.get("/")
def portal():
    html = render_template("portal.html")
    html = html.replace('<div>LOGO AB<small>ESPACIO RESERVADO</small></div>',
                        '<img src="/static/logo-ab-premiere.png" alt="AB Premiere" style="width:100%;height:100%;object-fit:cover;border-radius:12px">')
    return html.replace('placeholder="EJ. AB-2026-VIP"',
                        'placeholder="EJ. 2026" maxlength="4" pattern="[0-9]{4}" inputmode="numeric"')


@app.get("/admin")
def admin():
    html = render_template("admin.html")
    return html.replace('<a class="btn" href="/api/admin/export.csv">Exportar CSV</a>',
                        '<a class="btn" href="/api/admin/export.xlsx">Exportar Excel</a> <a class="btn gray" href="/api/admin/export.csv">CSV</a>')


@app.get("/api/photos/<path:name>")
def photo(name):
    from flask import send_from_directory
    return send_from_directory(data_dir() / "photos", Path(name).name)


@app.post("/api/register")
def register():
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name", "")).strip()
    phone = str(payload.get("phone", "")).strip()
    release = str(payload.get("release_code", ""))
    if len(name) < 3:
        return json_error("Escribe el nombre completo.")
    with db() as conn:
        stored = conn.execute("SELECT value FROM settings WHERE key='release_code'").fetchone()
        if not stored or not check_password_hash(stored["value"], release):
            return json_error("Código de liberación incorrecto.", 403)
        for _ in range(100):
            code = f"{secrets.randbelow(10000):04d}"
            if not conn.execute("SELECT 1 FROM clients WHERE code=?", (code,)).fetchone():
                break
        else:
            return json_error("No fue posible generar un código disponible. Intenta nuevamente.", 503)
        try:
            photo_name = save_photo(payload.get("photo", ""), code)
            conn.execute("INSERT INTO clients(code,full_name,phone,photo) VALUES(?,?,?,?)",
                         (code, name, phone, photo_name))
        except (ValueError, base64.binascii.Error) as exc:
            return json_error(str(exc))
    return jsonify(ok=True, code=code, name=name)


@app.post("/api/checkin")
def checkin():
    code = str((request.get_json(silent=True) or {}).get("code", "")).strip().upper()
    if len(code) != 4 or not code.isdigit():
        return json_error("El código debe contener exactamente 4 números.")
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with db() as conn:
        row = conn.execute("SELECT * FROM clients WHERE code=? AND active=1", (code,)).fetchone()
        if not row:
            return json_error("Código no encontrado o socio inactivo.", 404)
        last = conn.execute("SELECT visited_at FROM visits WHERE client_id=? ORDER BY id DESC LIMIT 1",
                            (row["id"],)).fetchone()
        duplicate = bool(last and last["visited_at"][:10] == now[:10])
        if not duplicate:
            conn.execute("INSERT INTO visits(client_id,visited_at) VALUES(?,?)", (row["id"], now))
        count = conn.execute("SELECT COUNT(*) n FROM visits WHERE client_id=?", (row["id"],)).fetchone()["n"]
    result = client_payload(row, count)
    result.update(ok=True, duplicate=duplicate)
    return jsonify(result)


@app.post("/api/admin/login")
def admin_login():
    payload = request.get_json(silent=True) or {}
    with db() as conn:
        row = conn.execute("SELECT * FROM admins WHERE username=? AND active=1",
                           (str(payload.get("username", "")).strip(),)).fetchone()
    if not row or not check_password_hash(row["password_hash"], str(payload.get("password", ""))):
        return json_error("Usuario o contraseña incorrectos.", 401)
    session["admin_id"] = row["id"]
    session["admin_name"] = row["username"]
    session["role"] = row["role"]
    with db() as conn:
        conn.execute("UPDATE admins SET last_login=CURRENT_TIMESTAMP WHERE id=?",(row["id"],))
    return jsonify(ok=True, username=row["username"])


@app.post("/api/admin/logout")
def admin_logout():
    session.clear()
    return jsonify(ok=True)


@app.get("/api/admin/session")
def admin_session():
    return jsonify(ok=True, authenticated=require_admin(), username=session.get("admin_name", ""), role=session.get("role", ""))


@app.get("/api/admin/dashboard")
def dashboard():
    if not require_admin():
        return json_error("Sesión requerida.", 401)
    with db() as conn:
        totals = conn.execute("""SELECT
            (SELECT COUNT(*) FROM clients WHERE active=1) clients,
            (SELECT COUNT(*) FROM visits WHERE date(visited_at)=date('now','localtime')) today,
            (SELECT COUNT(*) FROM visits) visits""").fetchone()
        frequent = conn.execute("""SELECT c.id,c.code,c.full_name,c.phone,c.photo,c.active,c.created_at,COUNT(v.id) visits,
            MAX(v.visited_at) last_visit FROM clients c LEFT JOIN visits v ON v.client_id=c.id
            GROUP BY c.id ORDER BY visits DESC,c.full_name LIMIT 100""").fetchall()
        days = conn.execute("""SELECT date(visited_at) day,COUNT(*) total FROM visits
            WHERE visited_at >= datetime('now','-29 days','localtime') GROUP BY date(visited_at) ORDER BY day""").fetchall()
        new_month = conn.execute("SELECT COUNT(*) n FROM clients WHERE strftime('%Y-%m',created_at)=strftime('%Y-%m','now','localtime')").fetchone()["n"]
    return jsonify(ok=True, totals={**dict(totals), "new_month": new_month}, clients=[dict(x) for x in frequent], days=[dict(x) for x in days])


@app.patch("/api/admin/clients/<int:client_id>")
def update_client(client_id):
    if not require_admin():
        return json_error("Sesión requerida.", 401)
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name", "")).strip()
    if len(name) < 3:
        return json_error("Nombre inválido.")
    with db() as conn:
        result = conn.execute("UPDATE clients SET full_name=?,phone=?,active=? WHERE id=?",
                     (name, str(payload.get("phone", "")).strip(), int(bool(payload.get("active", True))), client_id))
        if not result.rowcount:
            return json_error("Cliente no encontrado.", 404)
    audit("update", "client", client_id, {"name": name, "active": payload.get("active", True)})
    return jsonify(ok=True)


@app.get("/api/admin/clients/<int:client_id>/visits")
def client_visits(client_id):
    if not require_admin(): return json_error("Sesión requerida.", 401)
    with db() as conn:
        client=conn.execute("SELECT * FROM clients WHERE id=?",(client_id,)).fetchone()
        if not client:return json_error("Cliente no encontrado.",404)
        rows=conn.execute("SELECT * FROM visits WHERE client_id=? ORDER BY visited_at DESC",(client_id,)).fetchall()
    return jsonify(ok=True,client=dict(client),visits=[dict(x) for x in rows])


@app.post("/api/admin/clients/<int:client_id>/visits")
def manual_visit(client_id):
    if not require_admin(): return json_error("Sesión requerida.", 401)
    p=request.get_json(silent=True) or {}; visited=str(p.get("visited_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")).replace("T"," ")[:19]
    with db() as conn:
        if not conn.execute("SELECT 1 FROM clients WHERE id=?",(client_id,)).fetchone():return json_error("Cliente no encontrado.",404)
        cur=conn.execute("INSERT INTO visits(client_id,visited_at,notes,created_by,source) VALUES(?,?,?,?,?)",(client_id,visited,str(p.get("notes",""))[:500],session["admin_id"],"manual"))
    audit("create","visit",cur.lastrowid,f"client={client_id}"); return jsonify(ok=True)


@app.delete("/api/admin/visits/<int:visit_id>")
def delete_visit(visit_id):
    if not require_admin(): return json_error("Sesión requerida.", 401)
    with db() as conn: result=conn.execute("DELETE FROM visits WHERE id=?",(visit_id,))
    if not result.rowcount:return json_error("Visita no encontrada.",404)
    audit("delete","visit",visit_id); return jsonify(ok=True)


@app.get("/api/admin/users")
def users():
    if not require_admin(): return json_error("Sesión requerida.",401)
    with db() as conn: rows=conn.execute("SELECT id,username,full_name,role,active,created_at,last_login FROM admins ORDER BY username").fetchall()
    return jsonify(ok=True,users=[dict(x) for x in rows])


@app.post("/api/admin/users")
def create_user():
    if not require_admin() or session.get("role")!="admin":return json_error("Permiso de administrador requerido.",403)
    p=request.get_json(silent=True) or {}; username=str(p.get("username","")).strip(); password=str(p.get("password","")); role=str(p.get("role","consulta"))
    if len(username)<3 or len(password)<8:return json_error("Usuario mínimo 3 caracteres y contraseña mínimo 8.")
    if role not in ("admin","encargado","recepcion","consulta"):return json_error("Rol inválido.")
    try:
        with db() as conn: cur=conn.execute("INSERT INTO admins(username,password_hash,full_name,role) VALUES(?,?,?,?)",(username,generate_password_hash(password),str(p.get("full_name",""))[:100],role))
    except sqlite3.IntegrityError:return json_error("Ese usuario ya existe.",409)
    audit("create","admin",cur.lastrowid,role); return jsonify(ok=True)


@app.patch("/api/admin/users/<int:user_id>")
def update_user(user_id):
    if not require_admin() or session.get("role")!="admin":return json_error("Permiso de administrador requerido.",403)
    p=request.get_json(silent=True) or {}; role=str(p.get("role","consulta"))
    if role not in ("admin","encargado","recepcion","consulta"):return json_error("Rol inválido.")
    with db() as conn:
        conn.execute("UPDATE admins SET full_name=?,role=?,active=? WHERE id=?",(str(p.get("full_name",""))[:100],role,int(bool(p.get("active",True))),user_id))
        if p.get("password"):
            if len(str(p["password"]))<8:return json_error("Contraseña mínima: 8 caracteres.")
            conn.execute("UPDATE admins SET password_hash=? WHERE id=?",(generate_password_hash(str(p["password"])),user_id))
    audit("update","admin",user_id,role); return jsonify(ok=True)


@app.get("/api/admin/audit")
def audit_list():
    if not require_admin():return json_error("Sesión requerida.",401)
    with db() as conn: rows=conn.execute("SELECT a.*,COALESCE(u.username,'sistema') username FROM audit_log a LEFT JOIN admins u ON u.id=a.admin_id ORDER BY a.id DESC LIMIT 300").fetchall()
    return jsonify(ok=True,events=[dict(x) for x in rows])


@app.get("/api/admin/export.csv")
def export_csv():
    if not require_admin():
        return json_error("Sesión requerida.", 401)
    with db() as conn:
        rows = conn.execute("""SELECT c.code,c.full_name,c.phone,c.active,c.created_at,
            COUNT(v.id),MAX(v.visited_at) FROM clients c LEFT JOIN visits v ON v.client_id=c.id
            GROUP BY c.id ORDER BY c.full_name""").fetchall()
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["Código", "Nombre", "Teléfono", "Activo", "Registro", "Visitas", "Última visita"])
    writer.writerows([list(row) for row in rows])
    from flask import Response
    return Response("\ufeff" + output.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=clientes_ab_premiere.csv"})


@app.get("/api/admin/export.xlsx")
def export_excel():
    if not require_admin():
        return json_error("Sesión requerida.", 401)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with db() as conn:
        clients = conn.execute("""SELECT c.code,c.full_name,c.phone,c.active,c.created_at,
            COUNT(v.id) visitas,MAX(v.visited_at) ultima_visita
            FROM clients c LEFT JOIN visits v ON v.client_id=c.id
            GROUP BY c.id ORDER BY c.full_name""").fetchall()
        visits = conn.execute("""SELECT v.visited_at,c.code,c.full_name,v.source,v.notes
            FROM visits v JOIN clients c ON c.id=v.client_id ORDER BY v.visited_at DESC""").fetchall()
    wb = Workbook()
    gold, green, white = "C9A250", "173E2E", "FFFFFF"
    ws = wb.active; ws.title = "Clientes"
    ws.append(["Código", "Nombre", "Teléfono", "Estado", "Fecha de registro", "Visitas", "Última visita"])
    for row in clients:
        ws.append([row["code"], row["full_name"], row["phone"], "Activo" if row["active"] else "Inactivo",
                   row["created_at"], row["visitas"], row["ultima_visita"] or ""])
    wv = wb.create_sheet("Visitas")
    wv.append(["Fecha y hora", "Código", "Cliente", "Origen", "Notas"])
    for row in visits: wv.append(list(row))
    summary = wb.create_sheet("Resumen", 0)
    summary.append(["AB PREMIERE", "REPORTE DE CLIENTES Y VISITAS"])
    summary.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    summary.append(["Clientes", len(clients)])
    summary.append(["Visitas", len(visits)])
    for sheet in (ws, wv):
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 24
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=green); cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(horizontal="center")
        for col in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(r, col).value or "") for r in range(1, min(sheet.max_row, 300) + 1)]
            sheet.column_dimensions[get_column_letter(col)].width = min(max(len(x) for x in values) + 3, 38)
    summary.column_dimensions["A"].width=20; summary.column_dimensions["B"].width=40
    summary["A1"].fill=PatternFill("solid",fgColor=green); summary["B1"].fill=PatternFill("solid",fgColor=green)
    summary["A1"].font=Font(color=gold,bold=True,size=16); summary["B1"].font=Font(color=white,bold=True,size=13)
    stream=BytesIO(); wb.save(stream); stream.seek(0)
    from flask import send_file
    audit("export", "excel", details=f"{len(clients)} clientes, {len(visits)} visitas")
    return send_file(stream, as_attachment=True, download_name=f"AB_Premiere_{datetime.now():%Y%m%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/api/admin/backup")
def backup():
    if not require_admin():
        return json_error("Sesión requerida.", 401)
    folder = data_dir() / "backups"
    folder.mkdir(exist_ok=True)
    target = folder / f"ab_premiere_{datetime.now():%Y%m%d_%H%M%S}.db"
    source, destination = db(), sqlite3.connect(target)
    try:
        source.backup(destination)
    finally:
        destination.close(); source.close()
    for old in sorted(folder.glob("ab_premiere_*.db"), reverse=True)[30:]:
        old.unlink(missing_ok=True)
    return jsonify(ok=True, file=target.name, folder=str(folder))


@app.get("/api/admin/backups")
def backups():
    if not require_admin(): return json_error("Sesión requerida.", 401)
    folder=data_dir()/"backups"; folder.mkdir(exist_ok=True)
    items=[{"name":p.name,"size":p.stat().st_size,"modified":datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds")} for p in sorted(folder.glob("ab_premiere_*.db"),reverse=True)]
    return jsonify(ok=True,backups=items)


@app.post("/api/admin/settings")
def settings():
    if not require_admin():
        return json_error("Sesión requerida.", 401)
    payload = request.get_json(silent=True) or {}
    new_release = str(payload.get("release_code", ""))
    new_password = str(payload.get("password", ""))
    if new_release and len(new_release) < 6:
        return json_error("El código de liberación debe tener al menos 6 caracteres.")
    if new_password and len(new_password) < 8:
        return json_error("La contraseña debe tener al menos 8 caracteres.")
    with db() as conn:
        if new_release:
            conn.execute("UPDATE settings SET value=? WHERE key='release_code'", (generate_password_hash(new_release),))
        if new_password:
            conn.execute("UPDATE admins SET password_hash=? WHERE id=?", (generate_password_hash(new_password), session["admin_id"]))
    return jsonify(ok=True)


def run_server():
    serve(app, host=HOST, port=PORT, threads=8, channel_timeout=30)


if __name__ == "__main__":
    init_db()
    threading.Thread(target=run_server, daemon=True).start()
    time.sleep(.8)
    webview.create_window(APP_NAME, f"http://{HOST}:{PORT}", width=1280, height=820,
                          min_size=(960, 650), confirm_close=True, background_color="#121212")
    webview.start()
